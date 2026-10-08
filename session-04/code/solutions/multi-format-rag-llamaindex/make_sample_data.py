"""Writes the PDF, Excel and Word sample files into docs/. Everything in them is made up.

Run it only if you want to regenerate those files: uv run make_sample_data.py
"""
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

DOCS = Path("docs")
DOCS.mkdir(exist_ok=True)


# ---- PDF: a travel policy, three pages, no markdown headings to split on ----

def make_travel_policy_pdf():
    styles = getSampleStyleSheet()
    story = []

    def heading(text):
        story.append(Paragraph(text, styles["Heading2"]))

    def para(text):
        story.append(Paragraph(text, styles["BodyText"]))
        story.append(Spacer(1, 8))

    story.append(Paragraph("Acme Business Travel Policy", styles["Title"]))
    para("Department: Finance and Admin. Version 2025. Effective 1 April 2025.")

    heading("1. Purpose and Scope")
    para(
        "This policy explains how employees book, approve and claim for business travel. "
        "It applies to all grades, L1 to L5, for travel inside India and abroad. Travel that is "
        "paid for by a client is outside this policy unless the client contract says otherwise. "
        "The money limits that depend on grade, such as the hotel limit per night and the daily "
        "meal allowance, are kept in the Expense Limits workbook and not repeated here."
    )
    heading("2. Booking Flights and Trains")
    para(
        "All tickets must be booked through the company travel desk in the Admin portal. Tickets "
        "booked directly by the employee are reimbursed only if the travel desk could not offer "
        "the same journey, and a screenshot of the desk's reply must be attached to the claim. "
        "Book at least 7 days before the travel date, because fares rise sharply close to departure."
    )
    para(
        "Trains are the default choice for journeys under 6 hours, for example Ahmedabad to Mumbai. "
        "Air travel is allowed for longer journeys or when a train is not available on the required "
        "day. The class of travel depends on grade and is listed in the Expense Limits workbook. "
        "Business class is only available to L5, or to anyone on a flight longer than 6 hours when "
        "the department head has approved it in advance."
    )
    story.append(PageBreak())

    heading("3. Pre-Approval of Trips")
    para(
        "Every trip needs the reporting manager's approval in the Admin portal before booking. "
        "A trip needs the additional approval of the department head if it is longer than 3 nights "
        "or if the estimated cost is more than 25,000 INR. Pre-approval requests for these trips "
        "must be raised at least 5 working days before departure."
    )
    para(
        "International travel always needs the approval of the department head and the finance "
        "controller, whatever the cost. The traveller must also carry a valid passport and visa and "
        "register the trip with the Admin desk so that travel insurance can be arranged. Approval "
        "is not needed again for a change of dates, but a change of destination needs a new request."
    )
    heading("4. Hotels and Local Travel")
    para(
        "Hotels must be booked through the travel desk. The maximum rate per night depends on grade "
        "and is listed in the Expense Limits workbook. A stay above the limit needs the department "
        "head's approval before booking, and the extra amount is otherwise recovered from the "
        "employee. In cities where a company guest house is available, employees are expected to use "
        "it instead of a hotel."
    )
    para(
        "For local travel, employees can claim metro, bus, auto-rickshaw and app-cab fares up to the "
        "daily local travel limit for their grade. Car rental for a whole day needs the manager's "
        "approval. Personal vehicle use is reimbursed at 9 INR per kilometre for a car and 4 INR per "
        "kilometre for a two-wheeler, with a trip log attached."
    )
    story.append(PageBreak())

    heading("5. Claiming Travel Expenses")
    para(
        "A travel claim must be submitted in the Finance portal within 10 working days of returning "
        "from the trip. Claims submitted later are paid only with the approval of the finance "
        "controller. Every claim needs the original ticket or boarding pass, the hotel invoice, and "
        "receipts for any single expense above 500 INR. Meals are paid as a fixed daily allowance and "
        "need no receipts."
    )
    para(
        "Advance payments are available for trips longer than 3 nights. The advance can be up to 70 "
        "percent of the estimated cost and must be settled within 10 working days of the trip. An "
        "unsettled advance is deducted from the next month's salary."
    )
    heading("6. Non-Reimbursable Items")
    para(
        "The company does not reimburse minibar and alcohol charges, laundry for trips shorter than "
        "5 nights, personal entertainment, traffic fines, or the cost of upgrading a ticket or room "
        "above the entitlement of the employee's grade, unless the upgrade was approved in advance."
    )
    heading("7. Exceptions")
    para(
        "Exceptions to this policy can be approved by the finance controller in writing. The employee "
        "should explain why the policy could not be followed and attach any supporting email. An "
        "exception approved once does not set a precedent for later trips."
    )

    SimpleDocTemplate(str(DOCS / "travel_policy.pdf"), pagesize=A4, title="Acme Business Travel Policy").build(story)


# ---- Excel: two sheets of limits and entitlements by grade ----

def make_expense_limits_xlsx():
    wb = Workbook()

    travel = wb.active
    travel.title = "Travel Limits"
    travel.append(["Grade", "Hotel Limit Per Night (INR)", "Meal Allowance Per Day (INR)",
                   "Local Travel Per Day (INR)", "Flight Class"])
    for row in [
        ("L1", 3500, 800, 600, "Economy"),
        ("L2", 4500, 1000, 800, "Economy"),
        ("L3", 6000, 1500, 1200, "Economy"),
        ("L4", 8000, 2000, 1800, "Premium Economy"),
        ("L5", 12000, 2500, 2500, "Business"),
    ]:
        travel.append(row)

    leave = wb.create_sheet("Leave Entitlement")
    leave.append(["Grade", "Paid Leave Days Per Year", "Sick Leave Days Per Year",
                  "Casual Leave Days Per Year", "Max Carry Forward Days"])
    for row in [
        ("L1", 12, 8, 5, 6),
        ("L2", 14, 8, 5, 7),
        ("L3", 16, 8, 6, 8),
        ("L4", 18, 10, 6, 10),
        ("L5", 20, 10, 8, 12),
    ]:
        leave.append(row)

    for sheet in (travel, leave):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in "ABCDE":
            sheet.column_dimensions[column].width = 30

    wb.save(DOCS / "expense_limits.xlsx")


# ---- Word: an onboarding checklist with real heading styles ----

def make_onboarding_docx():
    doc = Document()
    doc.add_heading("New Joiner Onboarding Checklist", level=1)
    doc.add_paragraph("Department: HR. This checklist is shared with every new joiner and their manager.")

    doc.add_heading("Before Day One", level=2)
    for item in [
        "HR sends the offer documents and a list of papers to bring: ID proof, address proof, and the relieving letter from the previous employer.",
        "IT prepares a laptop and a company email account. The laptop is collected from the IT desk on the first morning.",
        "The manager nominates a buddy from the same team.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Day One", level=2)
    for item in [
        "Report to reception by 9:30 AM and collect the access card.",
        "Complete the HR induction, which covers the employee handbook, the code of conduct and the leave policy.",
        "Set up the laptop with the IT desk, including multi-factor authentication and the VPN.",
        "Have lunch with the buddy and the team.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("First Week", level=2)
    for item in [
        "Finish the mandatory online courses on information security and prevention of harassment within 5 working days.",
        "Meet the manager to agree the goals for the first 90 days.",
        "Submit bank details and nominee details in the HR portal so that the first salary is not delayed.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("First Month", level=2)
    for item in [
        "Add family members to the health insurance within 30 days of joining.",
        "Attend the monthly new joiner session with the HR head.",
        "Complete the probation goal-setting form. Probation lasts 6 months and is reviewed at 3 and 6 months.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Who to Contact", level=2)
    doc.add_paragraph(
        "For HR questions, write to the HR helpdesk. For laptop, email or access problems, raise a "
        "ticket with the IT desk. For travel and expense questions, contact Finance and Admin."
    )

    doc.save(DOCS / "onboarding_checklist.docx")


if __name__ == "__main__":
    make_travel_policy_pdf()
    make_expense_limits_xlsx()
    make_onboarding_docx()
    print("Wrote travel_policy.pdf, expense_limits.xlsx and onboarding_checklist.docx to docs/.")
